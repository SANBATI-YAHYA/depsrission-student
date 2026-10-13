from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_bullet_list(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_title_page(document: Document) -> None:
    title = document.add_heading("Project Architecture Presentation", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph("Student Depression Analysis Project")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph()
    document.add_heading("Project Summary", level=1)
    document.add_paragraph(
        "This repository presents a data-science project focused on understanding and "
        "predicting depression risk among students. It combines dataset exploration, "
        "feature preprocessing, SVM-based modeling, and visual reporting with Power BI."
    )

    add_bullet_list(
        document,
        [
            "Domain: Student mental-health analytics",
            "Main output: Depression risk classification (binary target 0/1)",
            "Audience: Analysts, educators, and student-support stakeholders",
        ],
    )


def add_problem_and_objectives_page(document: Document) -> None:
    document.add_page_break()
    document.add_heading("Problem Statement and Objectives", level=1)
    document.add_paragraph(
        "Student mental health is a critical issue. The project aims to identify major risk "
        "factors from student profiles and build a practical predictive model to support "
        "earlier intervention decisions."
    )

    document.add_heading("Key Objectives", level=2)
    add_bullet_list(
        document,
        [
            "Analyze student depression data and key behavioral/academic indicators",
            "Identify important risk signals linked to depression outcomes",
            "Train a predictive SVM classifier using prepared features",
            "Present insights clearly via notebook outputs and a Power BI report",
        ],
    )

    document.add_heading("Dataset Snapshot", level=2)
    add_bullet_list(
        document,
        [
            "Source file: ANALYST DEPRESSION STUDENT/data/Student Depression Dataset.csv",
            "Scale: 27,901 rows and 18 columns",
            "Target: Depression (values 0 or 1)",
            "Example features: Gender, Age, Academic Pressure, Sleep Duration, CGPA, Financial Stress",
        ],
    )


def add_architecture_page(document: Document) -> None:
    document.add_page_break()
    document.add_heading("High-Level Architecture and Components", level=1)
    document.add_paragraph(
        "The architecture is notebook-centric and organized around data, modeling, and "
        "reporting assets."
    )

    document.add_heading("Core Components", level=2)
    add_bullet_list(
        document,
        [
            "README.md: Project goals, structure, technologies, and usage guidance",
            "ANALYST DEPRESSION STUDENT/data/: Main CSV dataset and data dictionary notebook",
            "MODEL SVM CLASS_DEPRESSION/: Dictionary notebook and SVM modeling notebook",
            "ANALYST DEPRESSION STUDENT/rapport pbi/: Power BI dashboard file (.pbix)",
            "ANALYST DEPRESSION STUDENT/background/: Design assets used for visual materials",
        ],
    )

    document.add_heading("Simple Architecture Diagram", level=2)
    document.add_paragraph("Dataset CSV")
    document.add_paragraph("   ↓")
    document.add_paragraph("Data Dictionary + Feature Preparation (Jupyter notebooks)")
    document.add_paragraph("   ↓")
    document.add_paragraph("SVM Training and Prediction (scikit-learn)")
    document.add_paragraph("   ↓")
    document.add_paragraph("Insights and Communication (Power BI + project visuals)")


def add_workflow_and_stack_page(document: Document) -> None:
    document.add_page_break()
    document.add_heading("Main Workflow / Data Flow and Technology Stack", level=1)

    document.add_heading("Workflow / Data Flow", level=2)
    add_bullet_list(
        document,
        [
            "Load student dataset from CSV",
            "Inspect schema and target column",
            "Encode categorical fields and map sleep-duration categories",
            "Split data into train/test subsets",
            "Impute missing values and prepare model-ready features",
            "Train SVC model and evaluate predictions",
            "Share results in Power BI and project documentation",
        ],
    )

    document.add_heading("Technology Stack", level=2)
    add_bullet_list(
        document,
        [
            "Python and Jupyter Notebook for analysis and experimentation",
            "scikit-learn for machine-learning modeling (SVM/SVC)",
            "Pandas-oriented tabular processing in notebooks",
            "Power BI (.pbix) for interactive reporting",
            "PNG design assets for project communication",
        ],
    )


def add_conclusion_page(document: Document) -> None:
    document.add_page_break()
    document.add_heading("Conclusion, Strengths, and Future Improvements", level=1)

    document.add_heading("Conclusion", level=2)
    document.add_paragraph(
        "The repository delivers a complete baseline pipeline: data exploration, supervised "
        "classification, and report-ready outputs for student depression analysis."
    )

    document.add_heading("Current Strengths", level=2)
    add_bullet_list(
        document,
        [
            "Clear focus on a relevant social and educational challenge",
            "Practical dataset-to-model workflow with interpretable steps",
            "Combination of modeling notebooks and visual BI reporting",
            "Simple structure that is easy to explore and extend",
        ],
    )

    document.add_heading("Possible Future Improvements", level=2)
    add_bullet_list(
        document,
        [
            "Refactor notebooks into reusable Python modules/pipelines",
            "Add reproducible environment setup and dependency pinning",
            "Include model validation metrics tracking and comparison baselines",
            "Add explainability analysis (feature importance / SHAP-style interpretation)",
            "Document data quality checks and bias/ethics considerations",
        ],
    )


def main() -> None:
    repository_root = Path(__file__).resolve().parents[1]
    output_path = repository_root / "docs" / "project-architecture-presentation.docx"

    document = Document()
    add_title_page(document)
    add_problem_and_objectives_page(document)
    add_architecture_page(document)
    add_workflow_and_stack_page(document)
    add_conclusion_page(document)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    print(f"Created: {output_path}")


if __name__ == "__main__":
    main()
